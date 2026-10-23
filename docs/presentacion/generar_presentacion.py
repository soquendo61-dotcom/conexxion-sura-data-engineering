"""Genera la presentación ejecutiva de sustentación (python-pptx).

Uso: python docs/presentacion/generar_presentacion.py
Salida: docs/presentacion/Presentacion_Ejecutiva.pptx
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SALIDA = Path(__file__).with_name("Presentacion_Ejecutiva.pptx")

NAVY = RGBColor(0x0F, 0x2A, 0x44)
TEAL = RGBColor(0x1B, 0x99, 0x8B)
AMBER = RGBColor(0xF2, 0xA5, 0x41)
RED = RGBColor(0xC8, 0x4B, 0x31)
LIGHT = RGBColor(0xF3, 0xF6, 0xF9)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
TEXT = RGBColor(0x1F, 0x29, 0x33)
MUTED = RGBColor(0x5B, 0x6B, 0x7B)
LINE = RGBColor(0xD5, 0xDE, 0xE6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ICE = RGBColor(0xCF, 0xE3, 0xF0)
HEAD, BODY = "Cambria", "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


# ----------------------------------------------------------------------------- helpers
def fondo(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def texto(slide, x, y, w, h, contenido, size=16, color=TEXT, bold=False, font=BODY, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, italic=False, nombre=None):
    """contenido: str o lista de (texto, {opciones}) por párrafo."""
    caja = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if nombre:
        caja.name = nombre
    tf = caja.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    parrafos = contenido if isinstance(contenido, list) else [(contenido, {})]
    for i, (t, o) in enumerate(parrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = o.get("align", align)
        p.space_after = Pt(o.get("after", 0))
        r = p.add_run()
        r.text = t
        r.font.size = Pt(o.get("size", size))
        r.font.bold = o.get("bold", bold)
        r.font.italic = o.get("italic", italic)
        r.font.name = o.get("font", font)
        r.font.color.rgb = o.get("color", color)
    return caja


def bullets(slide, x, y, w, h, items, size=15, color=TEXT, after=6):
    caja = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = caja.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(after)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Inches(0.22)))
        pPr.set("indent", str(-Inches(0.22)))
        bu = pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", {"char": "•"})
        pPr.append(bu)
        if isinstance(it, tuple):
            r = p.add_run()
            r.text = it[0]
            r.font.bold = True
            r.font.size, r.font.name, r.font.color.rgb = Pt(size), BODY, color
            r2 = p.add_run()
            r2.text = it[1]
            r2.font.size, r2.font.name, r2.font.color.rgb = Pt(size), BODY, color
        else:
            r = p.add_run()
            r.text = it
            r.font.size, r.font.name, r.font.color.rgb = Pt(size), BODY, color
    return caja


def caja(slide, x, y, w, h, relleno=CARD, borde=None, forma=MSO_SHAPE.ROUNDED_RECTANGLE, radio=0.08, nombre=None):
    s = slide.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    if nombre:
        s.name = nombre
    s.fill.solid()
    s.fill.fore_color.rgb = relleno
    if borde is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = borde
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if forma == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radio
    s.text_frame.margin_left = s.text_frame.margin_right = Inches(0.1)
    return s


def circulo(slide, x, y, d, etiqueta, relleno=TEAL, color=WHITE, size=16):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid()
    s.fill.fore_color.rgb = relleno
    s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = etiqueta
    r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(size), True, BODY, color
    return s


def flecha(slide, x1, y1, x2, y2, color=MUTED):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(1.75)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd",
                             {"type": "triangle", "w": "med", "len": "med"}))
    return c


def titulo(slide, t, sub=None):
    texto(slide, 0.6, 0.45, 12.1, 0.8, t, size=32, bold=True, font=HEAD, color=NAVY, nombre="Titulo")
    if sub:
        texto(slide, 0.6, 1.2, 12.1, 0.45, sub, size=16, color=MUTED, italic=True, nombre="Subtitulo")


def pie(slide, n, oscuro=False):
    texto(slide, 0.6, 7.0, 9, 0.3, "Prueba técnica · Ingeniero de Datos · Pólizas y siniestros", size=10,
          color=ICE if oscuro else MUTED)
    texto(slide, 11.9, 7.0, 0.8, 0.3, str(n), size=10, color=ICE if oscuro else MUTED, align=PP_ALIGN.RIGHT)


def nueva(oscuro=False):
    s = prs.slides.add_slide(BLANK)
    fondo(s, NAVY if oscuro else LIGHT)
    return s


def tabla(slide, x, y, w, filas, anchos, alto_fila=0.4, size=13, encabezado=NAVY):
    t = slide.shapes.add_table(len(filas), len(filas[0]), Inches(x), Inches(y), Inches(w),
                               Inches(alto_fila * len(filas))).table
    for j, a in enumerate(anchos):
        t.columns[j].width = Inches(a)
    for i, fila in enumerate(filas):
        t.rows[i].height = Inches(alto_fila)
        for j, valor in enumerate(fila):
            celda = t.cell(i, j)
            celda.margin_left = celda.margin_right = Inches(0.08)
            celda.margin_top = celda.margin_bottom = Inches(0.03)
            celda.vertical_anchor = MSO_ANCHOR.MIDDLE
            celda.fill.solid()
            celda.fill.fore_color.rgb = encabezado if i == 0 else (CARD if i % 2 else LIGHT)
            tf = celda.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(valor)
            r.font.size, r.font.name = Pt(size), BODY
            r.font.bold = i == 0
            r.font.color.rgb = WHITE if i == 0 else TEXT
    return t


def notas(slide, t):
    slide.notes_slide.notes_text_frame.text = t


# ============================================================================ 1. Portada
s = nueva(oscuro=True)
for i, (cx, col) in enumerate([(9.6, TEAL), (10.9, AMBER), (12.2, ICE)]):
    circulo(s, cx - 0.5, 1.0 + i * 0.0, 1.0, ["B", "S", "G"][i], relleno=col, color=NAVY if col != TEAL else WHITE, size=26)
texto(s, 0.8, 2.3, 11.5, 1.8, "Pólizas y siniestros: de archivos operacionales a productos de datos certificados",
      size=40, bold=True, font=HEAD, color=WHITE)
texto(s, 0.8, 4.35, 11, 0.5, "Solución end-to-end en Databricks · Delta Lake · Unity Catalog · GitHub",
      size=20, color=ICE)
texto(s, 0.8, 5.6, 11, 0.9, [("Prueba técnica – Ingeniero de Datos", {"size": 16, "bold": True, "color": AMBER}),
                             ("Conexxion Sura · Octubre 2026", {"size": 14, "color": ICE})])
notas(s, "Presentación de la solución: cómo se llevaron cinco fuentes operacionales a productos de datos "
         "gobernados y certificados. Duración sugerida: 15-20 minutos + demo.")

# ============================================================================ 2. Resumen ejecutivo
s = nueva()
titulo(s, "Resumen: datos confiables, trazables y reprocesables",
       "Cinco fuentes integradas en una arquitectura Medallion con calidad medida y publicación certificada")
stats = [("5", "fuentes integradas\nBronze → Silver → Gold"), ("53", "registros en cuarentena\ncon causa y trazabilidad"),
         ("100%", "reconciliado\nconteos, claves y montos"), ("20", "pruebas automatizadas\n+ 13 controles por corrida")]
for i, (num, lab) in enumerate(stats):
    x = 0.6 + i * 3.08
    caja(s, x, 1.95, 2.85, 2.0)
    texto(s, x + 0.25, 2.1, 2.4, 0.9, num, size=44, bold=True, font=HEAD, color=TEAL if i != 1 else AMBER)
    texto(s, x + 0.25, 3.0, 2.4, 0.85, lab, size=14, color=MUTED)
bullets(s, 0.6, 4.35, 12.1, 2.5, [
    ("Qué resuelve: ", "estado integral de la póliza, prima emitida/recaudada/pendiente, siniestralidad mensual y discrepancias."),
    ("Cómo garantiza confianza: ", "data contracts, 60+ reglas con severidad y umbral, cuarentena, reconciliación entre capas."),
    ("Cómo se opera: ", "Workflow el día 5 a las 06:00 COT, idempotente, con bitácora, alertas y SLA de publicación."),
    ("Cómo escala: ", "lógica en módulos probados, Asset Bundles y promoción DEV → QA → PROD."),
], size=16, after=8)
pie(s, 2)
notas(s, "Mensaje principal: la solución no solo transforma datos, también garantiza que lo publicado es confiable "
         "y trazable, y que puede reprocesarse sin duplicar. Las cifras salen de la ejecución con los datos reales.")

# ============================================================================ 3. Problema
s = nueva()
titulo(s, "Punto de partida: fuentes con problemas reales",
       "Perfilamiento de los 5 datasets (19.512 filas) antes de diseñar")
problemas = [
    ("Duplicados", "Exactos y reenvíos del mismo hecho con otro ID"),
    ("Fechas", "3 formatos: yyyy-MM-dd, yyyy/MM/dd, dd/MM/yyyy"),
    ("Nulos", "Producto, estado, fecha de pago y de evento"),
    ("Huérfanos", "8 pólizas sin cliente, 22 hijos sin póliza"),
    ("Montos", "Primas, recaudos y pagos negativos"),
    ("Tardíos", "76% de recaudos ingeridos > 30 días después"),
    ("Esquema", "Columnas distintas al contrato esperado"),
    ("Empates", "Versiones con la misma marca temporal"),
    ("Incoherencias", "Fin de vigencia < inicio, aviso < ocurrencia"),
]
for i, (t, d) in enumerate(problemas):
    fila, col = divmod(i, 3)
    x, y = 0.6 + col * 4.1, 1.95 + fila * 1.6
    caja(s, x, y, 3.85, 1.38)
    circulo(s, x + 0.22, y + 0.34, 0.7, str(i + 1), relleno=NAVY if i % 2 else TEAL, size=16)
    texto(s, x + 1.12, y + 0.22, 2.6, 0.4, t, size=17, bold=True, color=NAVY)
    texto(s, x + 1.12, y + 0.62, 2.6, 0.7, d, size=13, color=MUTED)
pie(s, 3)
notas(s, "Todas las situaciones intencionales del enunciado se encontraron y cuantificaron. Decisión clave: no "
         "corregir arbitrariamente; separar errores (cuarentena) de excepciones válidas (marcadas y reportadas).")

# ============================================================================ 4. Arquitectura objetivo
s = nueva()
titulo(s, "Arquitectura objetivo: Lakehouse gobernado en Azure",
       "La prueba implementa el núcleo en un único workspace de Databricks Free Edition")
cols = [
    ("Fuentes", ["CRM", "Core pólizas / siniestros", "Pagos", "Bus de eventos"], MUTED),
    ("Ingesta", ["ADF / Lakeflow Connect", "CDC / watermark", "Event Hubs + Auto Loader"], TEAL),
    ("Lakehouse · ADLS + Delta", ["BRONZE  copia fiel", "SILVER  calidad + cuarentena", "GOLD  productos"], NAVY),
    ("Consumo", ["Vistas certificadas", "Power BI / SQL", "ML · GenAI (Genie)", "Apps operativas"], AMBER),
]
xs = [0.6, 3.75, 6.9, 10.05]
for (nombre, items, col), x in zip(cols, xs):
    caja(s, x, 1.85, 2.75, 3.45, relleno=CARD, borde=LINE)
    texto(s, x + 0.15, 1.97, 2.45, 0.4, nombre, size=15, bold=True, color=col if col != AMBER else RED)
    for k, it in enumerate(items):
        b = caja(s, x + 0.15, 2.5 + k * 0.68, 2.45, 0.55, relleno=col if col == NAVY else LIGHT, radio=0.2)
        tf = b.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = it
        r.font.size, r.font.name = Pt(12.5), BODY
        r.font.bold = col == NAVY
        r.font.color.rgb = WHITE if col == NAVY else TEXT
for x in xs[:-1]:
    flecha(s, x + 2.78, 3.55, x + 3.12, 3.55)
transv = [("Unity Catalog", "permisos · linaje · tags · máscaras PII"),
          ("Observabilidad", "ctl_* · SLA · alertas · system tables"),
          ("CI/CD", "GitHub Actions · Bundles · DEV→QA→PROD"),
          ("Seguridad", "Key Vault · Entra ID · service principals")]
for i, (t, d) in enumerate(transv):
    x = 0.6 + i * 3.08
    caja(s, x, 5.55, 2.88, 1.2, relleno=NAVY)
    texto(s, x + 0.2, 5.68, 2.6, 0.35, t, size=15, bold=True, color=AMBER)
    texto(s, x + 0.2, 6.05, 2.6, 0.6, d, size=12, color=WHITE)
pie(s, 4)
notas(s, "Arquitectura objetivo corporativa. En la prueba: volumen UC como landing, tablas Delta administradas, "
         "Workflow serverless y GitHub. Diagrama detallado en docs/architecture/diagrama_arquitectura.md.")

# ============================================================================ 5. Flujo implementado
s = nueva()
titulo(s, "Flujo implementado en cinco tareas",
       "Cada tarea es idempotente y registra estado, duración y volúmenes en ctl_ejecuciones")
pasos = [
    ("Bronze", "Hash SHA-256 por archivo, lectura como texto, metadatos de ingesta, contrato"),
    ("Silver", "Tipado multiformato, deduplicación, integridad, reglas, cuarentena, MERGE"),
    ("Gold", "poliza_360, siniestralidad mensual y discrepancias por fecha de corte"),
    ("Validar", "Reconciliación Silver→Gold, postcondiciones y frescura; publica el corte"),
    ("Gobierno", "Vistas certificadas, máscaras PII, tags y permisos en Unity Catalog"),
]
for i, (t, d) in enumerate(pasos):
    x = 0.6 + i * 2.42
    forma = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
    ch = caja(s, x, 1.95, 2.45, 0.95, relleno=[TEAL, NAVY, TEAL, NAVY, AMBER][i], forma=forma)
    tf = ch.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = f"{i + 1}. {t}"
    r.font.size, r.font.bold, r.font.name = Pt(17), True, BODY
    r.font.color.rgb = NAVY if i == 4 else WHITE
    caja(s, x + 0.05, 3.15, 2.3, 1.95, borde=LINE)
    texto(s, x + 0.22, 3.3, 1.98, 1.7, d, size=13.5, color=TEXT)
caja(s, 0.6, 5.4, 12.1, 1.35, relleno=CARD, borde=LINE)
texto(s, 0.85, 5.55, 11.6, 1.1, [
    ("Ejecución en un solo paso o como Workflow programado", {"size": 15, "bold": True, "color": NAVY, "after": 4}),
    ("Los notebooks solo orquestan: toda la lógica vive en src/ (módulos reutilizables, probados con pytest en local "
     "y en CI). Mismo run_id = {{job.run_id}} para todas las tareas de una ejecución.", {"size": 14, "color": MUTED}),
])
pie(s, 5)
notas(s, "Separación de responsabilidades Medallion. Un fallo en Silver por regla bloqueante impide que Gold se "
         "publique; los consumidores siguen viendo el último corte certificado.")

# ============================================================================ 6. Calidad
s = nueva()
titulo(s, "Calidad: cada regla tiene severidad, acción y umbral",
       "60+ reglas en 6 dimensiones; los errores se aíslan y las excepciones válidas se reportan")
sev = [("BLOQUEO", RED, "Supera el umbral → la tarea falla después de registrar evidencias; Gold no se publica",
        "Ej.: prima > 0 (umbral 2%)"),
       ("ADVERTENCIA", AMBER, "Supera el umbral → alerta al data steward; el pipeline continúa",
        "Ej.: póliza existe (2%), pago en vigencia (60%)"),
       ("INFORMATIVA", TEAL, "Solo métrica para tendencias y oportunidad",
        "Ej.: llegada tardía > 30 días")]
for i, (t, col, d, e) in enumerate(sev):
    y = 1.9 + i * 1.12
    caja(s, 0.6, y, 6.6, 0.98)
    circulo(s, 0.8, y + 0.19, 0.6, "!" if i == 0 else ("~" if i == 1 else "i"), relleno=col, size=18)
    texto(s, 1.6, y + 0.1, 5.4, 0.35, t, size=15, bold=True, color=NAVY)
    texto(s, 1.6, y + 0.43, 5.5, 0.5, f"{d}. {e}", size=12, color=MUTED)
caja(s, 0.6, 5.35, 6.6, 1.4, relleno=NAVY)
texto(s, 0.85, 5.48, 6.2, 1.2, [
    ("Acción por registro", {"size": 15, "bold": True, "color": AMBER, "after": 4}),
    ("CUARENTENA: sale de Silver con regla, fecha, run_id y registro original.", {"size": 13, "color": WHITE, "after": 2}),
    ("MARCAR: sigue a Silver con la regla en _dq_advertencias.", {"size": 13, "color": WHITE}),
])
cd = CategoryChartData()
cd.categories = ["clientes", "pólizas", "recaudos", "siniestros", "movimientos"]
cd.add_series("Duplicados", (2, 2, 2, 2, 2))
cd.add_series("Validez / consistencia", (0, 10, 14, 11, 8))
grafico = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(7.5), Inches(1.75), Inches(5.2), Inches(5.0), cd).chart
grafico.has_title = True
grafico.chart_title.text_frame.text = "Cuarentena por dataset (corte 2026-08-31)"
tp = grafico.chart_title.text_frame.paragraphs[0].runs[0].font
tp.size, tp.bold, tp.color.rgb, tp.name = Pt(13), True, NAVY, BODY
grafico.has_legend = True
grafico.legend.position = XL_LEGEND_POSITION.BOTTOM
grafico.legend.include_in_layout = False
grafico.legend.font.size, grafico.legend.font.name = Pt(11), BODY
for serie, color in zip(grafico.series, (AMBER, NAVY)):
    serie.format.fill.solid()
    serie.format.fill.fore_color.rgb = color
pl = grafico.plots[0]
pl.gap_width = 60
pl.has_data_labels = True
pl.data_labels.font.size, pl.data_labels.font.color.rgb = Pt(11), WHITE
pl.data_labels.position = XL_LABEL_POSITION.CENTER
pl.data_labels.number_format = '0;;;'
pl.data_labels.number_format_is_linked = False
grafico.value_axis.has_major_gridlines = False
grafico.value_axis.visible = False
grafico.category_axis.tick_labels.font.size = Pt(12)
grafico.category_axis.tick_labels.font.color.rgb = TEXT
grafico.category_axis.format.line.color.rgb = LINE
pie(s, 6)
notas(s, "Diferencia clave: un dato malo puntual va a cuarentena; un lote roto bloquea la publicación. "
         "Los huérfanos y pagos fuera de vigencia NO se rechazan: son discrepancias de negocio a investigar.")

# ============================================================================ 7. Productos de datos
s = nueva()
titulo(s, "Productos Gold orientados a las preguntas del negocio",
       "Una foto por fecha de corte: reprocesar un mes reemplaza solo ese mes")
prods = [
    ("gold_poliza_360", "1 fila por póliza y corte",
     "Estado al corte (vigente, próxima a vencer, vencida, renovada, cancelada), prima, recaudo, saldo, siniestros, última actividad"),
    ("gold_siniestralidad_mensual", "Mes × producto × sucursal",
     "Pólizas expuestas, prima devengada pro rata, emitida, recaudada, pendiente, costo incurrido y razón de siniestralidad"),
    ("gold_calidad_pipeline", "Control × dataset × ejecución",
     "Reglas, contratos, reconciliaciones, postcondiciones y frescura con resultado OK / ADVERTENCIA / FALLA"),
    ("gold_discrepancias", "Hallazgo × corte",
     "Pagos sin póliza, siniestros sin póliza válida, movimientos fuera de vigencia, estados inconsistentes"),
]
for i, (t, g, d) in enumerate(prods):
    fila, col = divmod(i, 2)
    x, y = 0.6 + col * 6.15, 1.9 + fila * 2.45
    caja(s, x, y, 5.95, 2.2)
    circulo(s, x + 0.25, y + 0.25, 0.65, ["360", "%", "✓", "≠"][i], relleno=[TEAL, NAVY, TEAL, AMBER][i],
            color=NAVY if i == 3 else WHITE, size=14)
    texto(s, x + 1.1, y + 0.22, 4.6, 0.4, t, size=17, bold=True, color=NAVY, font=BODY)
    texto(s, x + 1.1, y + 0.62, 4.6, 0.35, g, size=13, italic=True, color=TEAL)
    texto(s, x + 1.1, y + 1.02, 4.65, 1.1, d, size=13.5, color=TEXT)
pie(s, 7)
notas(s, "Granularidad, claves y métricas documentadas en docs/05_modelo_datos_y_calidad.md. Dimensiones "
         "conformadas: cliente, producto, sucursal, canal, periodo y fecha de corte. SCD2 en clientes.")

# ============================================================================ 8. Resultados
s = nueva()
titulo(s, "Resultados del corte 2026-08-31", "Ejecución con los datos reales de la prueba")
cd = CategoryChartData()
cd.categories = ["Vencida", "Cancelada", "Renovada", "Vigente", "Próxima a vencer"]
cd.add_series("Pólizas", (669, 284, 262, 236, 39))
g = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(1.8), Inches(6.4), Inches(4.9), cd).chart
g.has_title = True
g.chart_title.text_frame.text = "Pólizas por estado al corte (1.490)"
tp = g.chart_title.text_frame.paragraphs[0].runs[0].font
tp.size, tp.bold, tp.color.rgb, tp.name = Pt(13), True, NAVY, BODY
g.has_legend = False
pl = g.plots[0]
pl.gap_width = 50
pl.has_data_labels = True
pl.data_labels.font.size, pl.data_labels.font.color.rgb = Pt(12), TEXT
pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
ser = g.series[0]
ser.format.fill.solid()
ser.format.fill.fore_color.rgb = NAVY
for idx in (3, 4):
    pt = ser.points[idx]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = TEAL if idx == 3 else AMBER
g.value_axis.visible = False
g.value_axis.has_major_gridlines = False
g.category_axis.reverse_order = True
g.category_axis.tick_labels.font.size = Pt(12)
g.category_axis.format.line.color.rgb = LINE
kpis = [("$11.165 M", "prima emitida (COP)"), ("$7.212 M", "recaudo aplicado (COP)"),
        ("$79.523 M", "costo incurrido (COP)"), ("2.429 · 1.959", "pagos · siniestros fuera de vigencia")]
for i, (v, l) in enumerate(kpis):
    fila, col = divmod(i, 2)
    x, y = 7.35 + col * 2.75, 1.85 + fila * 1.55
    caja(s, x, y, 2.6, 1.35)
    texto(s, x + 0.18, y + 0.15, 2.3, 0.55, v, size=22, bold=True, font=HEAD, color=TEAL if i < 2 else (NAVY if i == 2 else RED))
    texto(s, x + 0.18, y + 0.75, 2.3, 0.5, l, size=12, color=MUTED)
caja(s, 7.35, 5.1, 5.35, 1.6, relleno=NAVY)
texto(s, 7.6, 5.22, 4.9, 1.4, [
    ("Lectura de negocio", {"size": 15, "bold": True, "color": AMBER, "after": 4}),
    ("Los costos muy altos y el gran volumen fuera de vigencia son propios del dato dummy: la solución los "
     "expone como discrepancias para investigación en lugar de ocultarlos.", {"size": 13, "color": WHITE}),
])
pie(s, 8)
notas(s, "Todas las reconciliaciones cuadran exactamente: 1.490 pólizas, prima 11.165.022.382,05, recaudo "
         "7.211.692.057,25 y costo 79.522.530.398,65 entre Silver, poliza_360 y siniestralidad.")

# ============================================================================ 9. Incremental e idempotencia
s = nueva()
titulo(s, "Reejecutar o reprocesar nunca duplica",
       "Idempotencia por capa y estrategia incremental para fuentes futuras")
tabla(s, 0.6, 1.9, 7.3, [
    ["Capa", "Mecanismo", "Clave"],
    ["Bronze", "replaceWhere + omitir archivos ya ingeridos", "hash SHA-256 del archivo"],
    ["Silver", "MERGE (solo si cambia el hash) + sincronización", "clave de negocio"],
    ["Cuarentena", "MERGE; RESUELTO si deja de fallar", "registro × regla"],
    ["Gold", "replaceWhere por periodo", "fecha_corte"],
    ["Control", "MERGE por tarea", "run_id × tarea"],
], [1.4, 3.6, 2.3], alto_fila=0.62, size=13)
evol = [("Archivos", "Hash + Auto Loader con checkpoint"), ("Bases de datos", "CDC por log o watermark con solape de 3 días"),
        ("APIs", "updated_since = último watermark exitoso"), ("Eventos", "Streaming + MERGE por evento_id")]
texto(s, 8.3, 1.9, 4.4, 0.4, "Evolución de fuentes", size=16, bold=True, color=NAVY)
for i, (t, d) in enumerate(evol):
    y = 2.4 + i * 0.86
    circulo(s, 8.3, y + 0.05, 0.55, str(i + 1), relleno=TEAL if i % 2 == 0 else NAVY, size=14)
    texto(s, 9.0, y, 3.7, 0.35, t, size=14, bold=True, color=TEXT)
    texto(s, 9.0, y + 0.35, 3.7, 0.5, d, size=12.5, color=MUTED)
caja(s, 0.6, 6.0, 12.1, 0.8, relleno=CARD, borde=LINE)
texto(s, 0.85, 6.12, 11.6, 0.6, "Backfill: fecha_corte = AAAA-MM-DD y modo = reproceso → recalcula y reemplaza solo esa foto. "
      "El watermark solo avanza tras una publicación exitosa.", size=14, color=TEXT)
pie(s, 9)
notas(s, "Demostrado en notebooks/09: la reejecución con los mismos archivos deja conteos idénticos. Llegadas "
         "tardías: entran en el siguiente proceso y el periodo afectado se corrige con backfill.")

# ============================================================================ 10. Orquestación y observabilidad
s = nueva()
titulo(s, "Operación: programada, observable y con SLA",
       "Día 5 de cada mes · 06:00 a. m. hora Colombia · cron 0 0 6 5 * ? (America/Bogota)")
hitos = [("06:00", "Inicio del job"), ("~06:15", "Bronze · Silver · Gold"), ("~06:20", "Validación y publicación"),
         ("09:00", "Límite del SLA")]
flecha(s, 0.9, 2.45, 12.3, 2.45, color=LINE)
for i, (h, d) in enumerate(hitos):
    x = 0.8 + i * 3.6
    circulo(s, x, 2.15, 0.6, "", relleno=AMBER if i == 3 else TEAL)
    ancho = 2.4 if i < 3 else 1.8
    texto(s, x - 0.2, 2.85, ancho, 0.4, h, size=16, bold=True, color=NAVY)
    texto(s, x - 0.2, 3.22, ancho, 0.4, d, size=13, color=MUTED)
ops = [("Reintentos y timeouts", "2 reintentos cada 5 min por tarea; Silver 1 (fallos de calidad son determinísticos)"),
       ("Alertas", "Correo on-failure y duración > 1 h; corporativo: Teams / PagerDuty"),
       ("Recuperación", "Repair run desde la tarea fallida; archivos RECHAZADO; RESTORE con time travel")]
for i, (t, d) in enumerate(ops):
    y = 4.0 + i * 0.95
    caja(s, 0.6, y, 6.1, 0.82)
    texto(s, 0.85, y + 0.08, 5.7, 0.3, t, size=14, bold=True, color=NAVY)
    texto(s, 0.85, y + 0.4, 5.7, 0.4, d, size=12, color=MUTED)
caja(s, 7.0, 4.0, 5.7, 2.72, relleno=NAVY)
texto(s, 7.25, 4.12, 5.3, 0.4, "Tablero de control (vw_obs_*)", size=15, bold=True, color=AMBER)
bullets(s, 7.25, 4.6, 5.3, 2.1, ["Estado y duración por tarea", "Volúmenes, aceptados, rechazados, duplicados",
                                  "Frescura y cumplimiento del SLA", "Evolución del esquema y costo estimado (DBU)",
                                  "Resultado de cada regla y reconciliación"], size=13, color=WHITE, after=3)
pie(s, 10)
notas(s, "Acuerdos: archivos a más tardar el día 4; corte certificado el día 5 a las 09:00; P1 con respuesta en 30 "
         "minutos. Procedimiento de atención documentado en docs/06_operacion_observabilidad.md.")

# ============================================================================ 11. Gobierno y seguridad
s = nueva()
titulo(s, "Gobierno con Unity Catalog y mínimo privilegio",
       "Todo objeto (volumen, tablas, vistas, funciones) registrado, documentado y con dueño")
tabla(s, 0.6, 1.9, 7.4, [
    ["Rol", "Bronze", "Silver", "Cuarentena", "Gold / vistas", "PII"],
    ["data_engineers", "✓", "✓", "✓", "✓", "masc."],
    ["data_stewards", "–", "✓", "✓", "✓", "masc."],
    ["data_analysts", "–", "–", "–", "✓", "–"],
    ["pii_lectores", "–", "✓", "–", "–", "✓"],
    ["Evaluadores", "lectura", "lectura", "lectura", "lectura", "masc."],
], [1.9, 1.0, 1.0, 1.2, 1.3, 1.0], alto_fila=0.55, size=13)
seg = [("Clasificación", "Documento, correo, teléfono = PII alta; nacimiento = media (tags)"),
       ("Enmascaramiento", "Máscaras de columna fn_mask_*: solo pii_lectores ve el dato"),
       ("Linaje y auditoría", "Linaje UC automático + _record_id, _file_hash, run_id por fila"),
       ("Datos personales", "Ley 1581: retención limitada, supresión con DELETE + VACUUM")]
for i, (t, d) in enumerate(seg):
    y = 1.9 + i * 1.2
    caja(s, 8.3, y, 4.4, 1.05)
    texto(s, 8.5, y + 0.1, 4.0, 0.3, t, size=14, bold=True, color=NAVY)
    texto(s, 8.5, y + 0.42, 4.05, 0.6, d, size=12, color=MUTED)
caja(s, 0.6, 5.45, 7.4, 1.3, relleno=NAVY)
texto(s, 0.85, 5.58, 7.0, 1.1, [
    ("Sin secretos en código", {"size": 15, "bold": True, "color": AMBER, "after": 4}),
    ("La solución no requiere credenciales; corporativamente Key Vault + secret scopes u OIDC. CI valida que no se "
     "versionen datos ni secretos.", {"size": 13, "color": WHITE}),
])
pie(s, 11)
notas(s, "Gold no contiene PII. Bronze y cuarentena (que guarda el registro original) son de acceso restringido. "
         "En Free Edition los grupos pueden no existir: esas sentencias se reportan como NO_APLICADA.")

# ============================================================================ 12. CI/CD y ambientes
s = nueva()
titulo(s, "Promoción del mismo artefacto: DEV → QA → PROD",
       "La implementación usa un único ambiente; la estrategia corporativa aísla workspaces y catálogos")
amb = [("DEV", "Fusión en dll → deploy automático", "Datos sintéticos / enmascarados", TEAL),
       ("QA", "Tag vX.Y.Z → deploy + ejecución de pruebas", "Subconjunto enmascarado de producción", NAVY),
       ("PROD", "Aprobación manual (environment protegido)", "Datos reales · run_as service principal", AMBER)]
for i, (t, d1, d2, col) in enumerate(amb):
    x = 0.6 + i * 4.2
    caja(s, x, 1.95, 3.7, 2.2)
    circulo(s, x + 0.25, 2.15, 0.85, t, relleno=col, color=NAVY if col == AMBER else WHITE, size=14)
    texto(s, x + 1.3, 2.2, 2.25, 0.8, d1, size=13.5, bold=True, color=TEXT)
    texto(s, x + 0.3, 3.2, 3.2, 0.8, d2, size=13, color=MUTED)
    if i < 2:
        flecha(s, x + 3.75, 3.05, x + 4.15, 3.05)
pasos_ci = [("Pull Request", "Revisión obligatoria hacia dll"), ("CI", "Ruff + 20 pruebas + validación del bundle"),
            ("Artefacto", "Asset Bundle versionado (commit / tag)"), ("Rollback", "Redeploy del tag anterior + RESTORE")]
for i, (t, d) in enumerate(pasos_ci):
    x = 0.6 + i * 3.08
    caja(s, x, 4.5, 2.88, 1.6, relleno=CARD, borde=LINE)
    circulo(s, x + 0.2, 4.68, 0.5, str(i + 1), relleno=NAVY, size=13)
    texto(s, x + 0.85, 4.72, 1.9, 0.4, t, size=14, bold=True, color=NAVY)
    texto(s, x + 0.2, 5.3, 2.5, 0.75, d, size=12.5, color=MUTED)
texto(s, 0.6, 6.35, 12.1, 0.4, "Catálogos por ambiente (dev_seguros, qa_seguros, prod_seguros) con workspace binding: DEV no puede leer PROD.",
      size=13, italic=True, color=TEXT)
pie(s, 12)
notas(s, "Trade-off: tres workspaces cuestan más y exigen sincronizar configuración, pero dan aislamiento fuerte, "
         "esperado en el sector asegurador. Definición en databricks.yml y .github/workflows.")

# ============================================================================ 13. Riesgos y evolución
s = nueva(oscuro=True)
texto(s, 0.6, 0.45, 12.1, 0.8, "Riesgos asumidos y líneas de evolución", size=32, bold=True, font=HEAD, color=WHITE)
texto(s, 0.6, 1.2, 12, 0.45, "Decisiones explícitas y próximos pasos hacia una solución corporativa", size=16,
      italic=True, color=ICE)
riesgos = ["Fechas dd/MM vs MM/dd: supuesto día/mes, ratificar en el contrato",
           "Estado de póliza sin historia: backfill usa el estado vigente",
           "Fecha efectiva de cancelación inferida de fecha_actualizacion",
           "Umbrales calibrados con datos dummy: acordarlos con los owners",
           "Free Edition: un ambiente, serverless, grupos limitados"]
evolucion = ["Silver incremental con Change Data Feed",
             "SCD2 de pólizas desde eventos para cortes as-of exactos",
             "Reglas y umbrales externalizados (DQX / expectativas)",
             "AI/BI Dashboard y alertas SQL; costos con system.billing",
             "Streaming de eventos y Genie space sobre vistas certificadas"]
for col_i, (tit, items, color) in enumerate([("Riesgos y supuestos", riesgos, AMBER), ("Evolución", evolucion, TEAL)]):
    x = 0.6 + col_i * 6.2
    caja(s, x, 1.95, 5.95, 3.75, relleno=RGBColor(0x17, 0x3A, 0x5C))
    texto(s, x + 0.3, 2.12, 5.4, 0.45, tit, size=19, bold=True, color=color)
    bullets(s, x + 0.3, 2.75, 5.4, 3.8, items, size=15, color=WHITE, after=10)
for i, palabra in enumerate(["Reproducible", "Auditable", "Reutilizable"]):
    texto(s, 0.6 + i * 4.1, 5.95, 3.9, 0.6, palabra, size=26, bold=True, font=HEAD,
          color=[TEAL, AMBER, ICE][i], align=PP_ALIGN.CENTER)
pie(s, 13, oscuro=True)
notas(s, "Cierre: la solución es reproducible (Git + bundle), auditable (linaje, control, time travel) y reutilizable "
         "(contratos, motor de reglas y tareas parametrizadas sirven para nuevos dominios). Pasar a la demo.")

prs.save(SALIDA)
print(f"Presentación generada: {SALIDA}")
